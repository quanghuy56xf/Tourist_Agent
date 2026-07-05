import io
import re
from typing import BinaryIO

from app.modules.rag.types import StructuredSection

MARKDOWN_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)$")
DOCX_HEADING_RE = re.compile(r"^Heading\s+(\d+)$", re.IGNORECASE)
HISTORICAL_CHAPTER_RE = re.compile(
    r"^Chương\s+(?:[IVXLCDM]+|\d+|[A-ZÀ-Ỵ][\wÀ-ỹ-]*)(?:\b.*)?$",
    re.IGNORECASE,
)
HISTORICAL_ROMAN_RE = re.compile(
    r"^[IVXLCDM]{1,8}\s*[-–—.]\s*\S.*$",
    re.IGNORECASE,
)
HISTORICAL_NUMBER_RE = re.compile(r"^\d{1,2}\s*[-–—.]\s*\S.*$")


def _flush_section(
    sections: list[StructuredSection],
    heading: str | None,
    level: int | None,
    body_lines: list[str],
    source_hint: str | None,
) -> None:
    body = "\n".join(line.rstrip() for line in body_lines).strip()
    if not heading and not body:
        return
    sections.append(
        StructuredSection(
            heading=heading,
            heading_level=level,
            body=body,
            source_hint=source_hint,
        )
    )


def _has_markdown_headings(text: str) -> bool:
    return any(
        MARKDOWN_HEADING_RE.match(line.strip())
        for line in text.split("\n")
        if line.strip()
    )


def _match_historical_heading(line: str) -> tuple[int, str] | None:
    stripped = line.strip()
    if not stripped:
        return None
    if HISTORICAL_CHAPTER_RE.match(stripped):
        return 1, stripped
    if HISTORICAL_ROMAN_RE.match(stripped):
        return 2, stripped
    if HISTORICAL_NUMBER_RE.match(stripped):
        return 3, stripped
    return None


def _is_part_heading(line: str) -> bool:
    return bool(re.match(r"^PHẦN\s+", line.strip(), re.IGNORECASE))


def _is_major_historical_heading(line: str) -> bool:
    stripped = line.strip()
    return bool(_is_part_heading(stripped) or HISTORICAL_CHAPTER_RE.match(stripped))


def _is_toc_like_line(line: str) -> bool:
    stripped = line.strip()
    if not stripped:
        return False
    return bool(
        _is_part_heading(stripped)
        or _match_historical_heading(stripped)
        or re.match(r"^[.•]\s*\S+", stripped)
    )


def _find_toc_start(lines: list[str]) -> int | None:
    scan_limit = min(len(lines), 200)
    window_size = 30
    min_heading_count = 6

    for start in range(scan_limit):
        window = [line for line in lines[start : start + window_size] if line.strip()]
        if len(window) < min_heading_count:
            continue
        toc_like_count = sum(1 for line in window if _is_toc_like_line(line))
        long_body_count = sum(1 for line in window if len(line.strip()) > 120)
        if toc_like_count >= min_heading_count and toc_like_count / len(window) >= 0.45 and long_body_count <= 2:
            return start
    return None


def _drop_table_of_contents(text: str) -> str:
    lines = text.split("\n")
    toc_start = _find_toc_start(lines)
    if toc_start is None:
        return text

    seen_major_headings: set[str] = set()
    for index in range(toc_start, len(lines)):
        normalized = lines[index].strip().casefold()
        if not normalized:
            continue
        if _is_major_historical_heading(lines[index]):
            if normalized in seen_major_headings:
                return "\n".join(lines[:toc_start] + lines[index:]).strip()
            seen_major_headings.add(normalized)

    return text


def _has_historical_headings(text: str) -> bool:
    chapter_count = 0
    roman_count = 0
    number_count = 0

    for line in text.split("\n"):
        match = _match_historical_heading(line)
        if match is None:
            continue
        level, _title = match
        if level == 1:
            chapter_count += 1
        elif level == 2:
            roman_count += 1
        elif level == 3:
            number_count += 1

    structured_count = roman_count + number_count
    return (chapter_count >= 1 and structured_count >= 1) or roman_count >= 2


def parse_text_sections(text: str) -> list[StructuredSection]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not normalized:
        return []
    normalized = _drop_table_of_contents(normalized)

    if _has_markdown_headings(normalized):
        return _parse_markdown_sections(normalized)

    if _has_historical_headings(normalized):
        historical = _parse_historical_sections(normalized)
        if historical:
            return historical

    heuristic = _parse_heuristic_sections(normalized)
    if heuristic:
        return heuristic

    return [
        StructuredSection(
            heading=None,
            heading_level=None,
            body=normalized,
            source_hint="text:flat",
        )
    ]


def _parse_markdown_sections(text: str) -> list[StructuredSection]:
    sections: list[StructuredSection] = []
    heading: str | None = None
    level: int | None = None
    body_lines: list[str] = []

    for line in text.split("\n"):
        match = MARKDOWN_HEADING_RE.match(line.strip())
        if match:
            _flush_section(sections, heading, level, body_lines, "md")
            heading = match.group(2).strip()
            level = len(match.group(1))
            body_lines = []
            continue
        body_lines.append(line)

    _flush_section(sections, heading, level, body_lines, "md")
    return sections


def _parse_historical_sections(text: str) -> list[StructuredSection]:
    sections: list[StructuredSection] = []
    context: list[str] = []
    body_lines: list[str] = []

    def flush_body() -> None:
        body = "\n".join(line.rstrip() for line in body_lines).strip()
        body_lines.clear()
        if not body:
            return
        heading = f"[Ngữ cảnh: {' > '.join(context)}]" if context else None
        sections.append(
            StructuredSection(
                heading=heading,
                heading_level=len(context) if context else None,
                body=body,
                source_hint="historical",
            )
        )

    for block in re.split(r"\n\s*\n+", text):
        lines = [line.strip() for line in block.split("\n") if line.strip()]
        if not lines:
            continue

        for line in lines:
            match = _match_historical_heading(line)
            if match is not None:
                flush_body()
                level, title = match
                if level == 1:
                    context = [title]
                elif level == 2:
                    context = context[:1] + [title]
                else:
                    context = context[:2] + [title]
                continue
            body_lines.append(line)

        flush_body()

    return sections


def _parse_heuristic_sections(text: str) -> list[StructuredSection]:
    blocks = re.split(r"\n\s*\n", text)
    if len(blocks) <= 1:
        return []

    sections: list[StructuredSection] = []
    for block in blocks:
        lines = [line.strip() for line in block.split("\n") if line.strip()]
        if not lines:
            continue
        first = lines[0]
        if (
            len(lines) > 1
            and len(first) < 80
            and not first.endswith(".")
        ):
            sections.append(
                StructuredSection(
                    heading=first,
                    heading_level=2,
                    body="\n".join(lines[1:]),
                    source_hint="heuristic",
                )
            )
        else:
            sections.append(
                StructuredSection(
                    heading=None,
                    heading_level=None,
                    body=block.strip(),
                    source_hint="heuristic",
                )
            )
    return sections if len(sections) > 1 else []


def parse_docx_sections(content: bytes) -> list[StructuredSection]:
    from docx import Document

    document = Document(io.BytesIO(content))
    full_text = "\n\n".join(
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    )
    full_text = _drop_table_of_contents(full_text)
    if _has_historical_headings(full_text):
        historical = _parse_historical_sections(full_text)
        if historical:
            return historical

    sections: list[StructuredSection] = []
    heading: str | None = None
    level: int | None = None
    body_lines: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name or "").strip()
        heading_match = DOCX_HEADING_RE.match(style_name)
        if heading_match:
            _flush_section(
                sections,
                heading,
                level,
                body_lines,
                f"docx:{style_name}",
            )
            heading = text
            level = int(heading_match.group(1))
            body_lines = []
            continue
        body_lines.append(text)

    _flush_section(sections, heading, level, body_lines, "docx")
    if sections:
        return sections

    return parse_text_sections(full_text)


def parse_flat_text_sections(
    text: str,
    *,
    source_hint: str = "text:flat",
) -> list[StructuredSection]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not normalized:
        return []
    normalized = _drop_table_of_contents(normalized)

    if _has_markdown_headings(normalized):
        return _parse_markdown_sections(normalized)

    if _has_historical_headings(normalized):
        historical = _parse_historical_sections(normalized)
        if historical:
            return historical

    return [
        StructuredSection(
            heading=None,
            heading_level=None,
            body=normalized,
            source_hint=source_hint,
        )
    ]


def extract_pdf_text(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    page_texts: list[str] = []

    for page in reader.pages:
        page_text = (page.extract_text() or "").strip()
        if page_text:
            page_texts.append(page_text)

    return "\n\n".join(page_texts).strip()


def parse_pdf_sections(content: bytes) -> list[StructuredSection]:
    text = extract_pdf_text(content)
    if not text:
        return []
    return parse_flat_text_sections(text, source_hint="pdf:txt")


def parse_upload_to_sections(
    *,
    source_type: str,
    content: bytes | str,
    filename: str | None = None,
) -> list[StructuredSection]:
    if source_type == "text":
        text = content if isinstance(content, str) else content.decode("utf-8")
        return parse_text_sections(text)

    if source_type == "txt":
        text = content.decode("utf-8") if isinstance(content, bytes) else content
        return parse_text_sections(text)

    if source_type == "docx":
        if isinstance(content, str):
            raise ValueError("DOCX content must be bytes")
        return parse_docx_sections(content)

    if source_type == "pdf":
        if isinstance(content, str):
            raise ValueError("PDF content must be bytes")
        return parse_pdf_sections(content)

    raise ValueError(f"Unsupported source type: {source_type}")


def sections_to_plain_text(sections: list[StructuredSection]) -> str:
    parts: list[str] = []
    for section in sections:
        if section.heading:
            parts.append(section.heading)
        if section.body:
            parts.append(section.body)
    return "\n\n".join(parts).strip()


def detect_source_type(filename: str | None, mime_type: str | None) -> str | None:
    if filename:
        lower = filename.lower()
        if lower.endswith(".txt"):
            return "txt"
        if lower.endswith(".docx"):
            return "docx"
        if lower.endswith(".pdf"):
            return "pdf"
    if mime_type == "text/plain":
        return "txt"
    if mime_type == "application/pdf":
        return "pdf"
    if mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return "docx"
    return None
